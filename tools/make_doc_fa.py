"""
make_doc_fa.py -- the project documentation in Persian as Word: docs/VOID_RUNNER_documentation_fa.docx
(right to left; open it in Google Docs: Drive -> upload, then "Open with Google Docs").

Same sections, pictures and tables as tools/make_doc.py: it runs make_doc's section functions
with a writer that translates every text (tools/fa_text.json, shared with make_latex_fa.py)
and marks every Persian paragraph and table as right to left. The code stays English, left to
right.

Usage (from the project folder, after `python tests/host/run_tests.py` built build/sim.exe):
    python tools/make_doc_fa.py
"""
import subprocess
import sys
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'tools'))
import make_doc as md                                        # noqa: E402
import make_latex_fa as fa                                   # noqa: E402  (translations)
import render_dumps as rd                                    # noqa: E402

OUT = ROOT / 'docs' / 'VOID_RUNNER_documentation_fa.docx'
FONT = 'Tahoma'                                              # has Persian letters on Windows
T = fa.translate


# The main menu with each item chosen (only in the Persian document: its own part, 3.6).
MENU_SCRIPT = '''seed 7
until "[VR] screen title" 10000
tap R
until "[VR] screen menu"
dump
tap S
dump
tap S
dump
tap S
dump
'''


def menu_pictures():
    script = md.BUILD / 'doc_menu.script'
    script.write_text(MENU_SCRIPT, encoding='utf-8', newline='\n')
    sim = md.BUILD / ('sim.exe' if sys.platform == 'win32' else 'sim')
    log = subprocess.run([str(sim), str(script)], cwd=md.ROOT, capture_output=True, text=True,
                         errors='replace').stdout
    dumps = rd.read_dumps(log)
    assert len(dumps) == 4, 'got %d menu dumps' % len(dumps)
    return [md.screen_png(d, 'menu_%d' % i) for i, d in enumerate(dumps)]


def menu_section(doc):
    """The main menu in detail: what each item opens and how the keys work."""
    doc.h('منوی اصلی', 3)
    doc.rich('منوی اصلی پس از صفحهٔ عنوان با فشار **R** باز می‌شود و مرکز همهٔ بخش‌های دستگاه است. '
             'چهار گزینه دارد که هر کدام در یکی از نوارهای رنگی صفحه نوشته شده‌اند: **PLAY** در '
             'منطقهٔ قرمز بالا، **SHIP** و **OPTS** در منطقهٔ سفید وسط و **BEST** در منطقهٔ سبز '
             'پایین. هر نام چهار حرف کوچک است تا کامل در عرض ۱۶ پیکسلی صفحه جا شود. گزینهٔ '
             'انتخاب‌شده با یک خط زیر آن مشخص می‌شود و نوارهای وضعیت در منو خاموش‌اند تا فقط '
             'خود منو دیده شود.')
    doc.gallery([(p, l) for p, l in zip(doc.menu_pics, ['PLAY انتخاب‌شده', 'SHIP انتخاب‌شده',
                                                         'OPTS انتخاب‌شده', 'BEST انتخاب‌شده'])],
                4, 3, 'منوی اصلی با هر یک از چهار گزینه؛ خط زیر، گزینهٔ انتخاب‌شده را نشان می‌دهد.')
    doc.table(['گزینه', 'با R چه چیزی باز می‌شود', 'توضیح'], [
        ['PLAY', 'یک بازی تازه', 'بازی از مرحلهٔ ۱ با سفینه و سختی انتخاب‌شده شروع می‌شود؛ '
                                 'جان‌ها و مهمات طبق سختی و سفینه پر می‌شوند و امتیاز صفر است.'],
        ['SHIP', 'آشیانه', 'انتخاب یکی از سه سفینهٔ BOLT، DART و TWIN (بخش ۳.۷).'],
        ['OPTS', 'تنظیمات', 'روشنایی، صدا، سطح سختی و آزمون LED (بخش ۳.۷ و ۳.۱۳).'],
        ['BEST', 'بهترین امتیاز', 'بالاترین امتیاز، مرحله‌ای که به آن رسید و سختی آن (بخش ۳.۱۱).'],
    ], [2.5, 4, 10.5], 'گزینه‌های منوی اصلی.')
    doc.table(['کلید', 'کار در منو'], [
        ['W', 'رفتن به گزینهٔ بالایی (از PLAY به BEST می‌چرخد) با یک بوق کوتاه'],
        ['S', 'رفتن به گزینهٔ پایینی (از BEST به PLAY می‌چرخد) با یک بوق کوتاه'],
        ['R', 'باز کردن گزینهٔ انتخاب‌شده با یک بوق بلندتر'],
        ['Q', 'بازگشت به صفحهٔ عنوان'],
        ['M / B', 'روشن و خاموش کردن صدا و تغییر روشنایی (در همهٔ صفحه‌ها کار می‌کنند)'],
    ], [3, 14], 'کلیدها در منوی اصلی.')
    doc.rich('منو آخرین گزینهٔ انتخاب‌شده را به خاطر می‌سپارد: مثلاً پس از بازگشت از آشیانه، '
             'خط زیر هنوز روی SHIP است. پس از پایان یک بازی (GAME OVER یا پیروزی) یا خروج از بازی '
             'با Q در صفحهٔ مکث هم دوباره به همین منو برمی‌گردید. در کد، این بخش در فایل `menu.h` '
             'است (بخش ۴.۱۳): تابع `menuUpdate()` کلیدها را می‌خواند و گزینهٔ انتخاب‌شده را '
             'برمی‌گرداند و `menuDraw()` منو را می‌کشد؛ ماشین حالت `void_runner.ino` سپس صفحهٔ '
             'مربوط را باز می‌کند.')


class FaDoc(md.Doc):
    FIG_LABEL, TAB_LABEL = 'شکل', 'جدول'
    menu_pics = []

    def h(self, text, level):
        if text == '3.7 Hangar and options':        # the menu part goes at the end of 3.6
            menu_section(self)
        return super().h(T(text), level)

    def p(self, text='', **kw):
        return super().p(T(text), **kw)

    def rich(self, text):
        return super().rich(T(text))

    def bullets(self, items):
        super().bullets([T(i) for i in items])

    def picture(self, path, width_cm, caption=None):
        super().picture(path, width_cm, T(caption) if caption else None)

    def gallery(self, items, per_row, width_cm, caption=None):
        super().gallery([(p, T(l)) for p, l in items], per_row, width_cm,
                        T(caption) if caption else None)

    def table(self, header, rows, widths=None, caption=None, size=9.5):
        return super().table([T(h) for h in header], [[T(c) for c in r] for r in rows],
                             widths, T(caption) if caption else None, size)


def set_rtl(doc):
    """Marks every paragraph with Persian text as right to left (and its Persian runs), and
    every table as right to left. Code lines (Consolas) and pure-Latin cells stay left to right."""
    def para_rtl(par):
        text = par.text
        if not fa.PERSIAN.search(text):
            return
        ppr = par._p.get_or_add_pPr()
        if ppr.find(qn('w:bidi')) is None:
            ppr.insert(0, OxmlElement('w:bidi'))
        for run in par.runs:
            rpr = run._r.get_or_add_rPr()
            fonts = rpr.find(qn('w:rFonts'))
            if fonts is None:
                fonts = OxmlElement('w:rFonts')
                rpr.insert(0, fonts)
            fonts.set(qn('w:cs'), FONT)
            if fa.PERSIAN.search(run.text):
                rpr.append(OxmlElement('w:rtl'))
            if run.font.size:                                # Persian letters use the cs size
                szcs = OxmlElement('w:szCs')
                szcs.set(qn('w:val'), str(int(run.font.size.pt * 2)))
                rpr.append(szcs)
            if run.bold:
                rpr.append(OxmlElement('w:bCs'))

    for par in doc.paragraphs:
        para_rtl(par)
    for table in doc.tables:
        tblpr = table._tbl.tblPr
        tblpr.append(OxmlElement('w:bidiVisual'))            # first column on the right
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    para_rtl(par)
    # Persian in the styles too (headings, Normal): complex-script font and size
    for style in doc.styles:
        if style.type != 1 or style.font is None:
            continue
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            rpr.insert(0, fonts)
        fonts.set(qn('w:cs'), FONT)
        if style.font.size:
            szcs = OxmlElement('w:szCs')
            szcs.set(qn('w:val'), str(int(style.font.size.pt * 2)))
            rpr.append(szcs)


def cover(doc):
    c = fa.COVER_FA
    for _ in range(3):
        doc.p('')
    doc.p(c['university'], size=16, align=WD_ALIGN_PARAGRAPH.CENTER).runs[-1].bold = True
    doc.p(c['faculty'], size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.p('')
    doc.p(c['course'], size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.p(c['project'], size=12, align=WD_ALIGN_PARAGRAPH.CENTER, colour=md.GREY)
    for _ in range(2):
        doc.p('')
    t = doc.p('VOID RUNNER', size=40, align=WD_ALIGN_PARAGRAPH.CENTER)
    t.runs[-1].bold = True
    t.runs[-1].font.color.rgb = md.ACCENT
    doc.p('یک بازی تیراندازی فضایی عمودی روی Arduino Mega', size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.p('۴۳۲ LED تکی · صفحه‌کلید ۹ دکمه‌ای · ۴ بازر · بدون ماژول و کتابخانه', size=10.5,
          align=WD_ALIGN_PARAGRAPH.CENTER, colour=md.GREY)
    for _ in range(3):
        doc.p('')
    info = [('دانشجو', c['student']), ('شمارهٔ دانشجویی', c['student_id']), ('استاد', c['teacher']),
            ('درس', c['course']), ('نیم‌سال', c['semester']), ('تاریخ', c['date']),
            ('بستر', 'Arduino Mega 2560، شبیه‌سازی‌شده در Wokwi'), ('کد منبع', md.COVER['repo'])]
    t = doc.d.add_table(rows=len(info), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, (k, v) in enumerate(info):
        a, b = t.cell(r, 0), t.cell(r, 1)
        a.width, b.width = Cm(4), Cm(10)
        ra = a.paragraphs[0].add_run(k)
        ra.bold = True
        ra.font.size = Pt(11)
        ra.font.color.rgb = md.ACCENT
        b.paragraphs[0].add_run(v).font.size = Pt(11)
    doc.page_break()


def main():
    pics = md.make_media()
    doc = FaDoc()
    doc.menu_pics = menu_pictures()
    cover(doc)
    md.hardware_page(doc, pics)
    md.performance(doc, pics)
    md.code_part(doc, pics)
    set_rtl(doc.d)
    doc.d.save(OUT)
    print('wrote', OUT)
    if fa.missing:
        print('%d texts have no Persian translation (left in English):' % len(fa.missing))
        for t in fa.missing:
            print('  ', t[:100])


if __name__ == '__main__':
    main()
