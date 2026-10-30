"""
make_doc.py -- builds the project documentation docs/VOID_RUNNER_documentation.docx
(open it in Google Docs: Drive -> New -> File upload, then "Open with Google Docs").

Parts: 1 cover page, 2 the hardware (a picture of the whole circuit), 3 project performance
(how the device works, with pictures of every screen), 4 the full code with explanations.

All pictures come from the real code: the PC harness plays the game (tools/doc_record.script)
and its screen dumps are drawn in the real LED colours (render_dumps.py); the circuit is
drawn with Wokwi's own part drawings (render_circuit.py, needs Edge and the internet).

Usage (from the project folder, after `python tests/host/run_tests.py` built build/sim.exe):
    python tools/make_doc.py
Needs python-docx and Pillow (pip install python-docx pillow).
"""
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'tools'))
import render_dumps as rd                                    # noqa: E402

BUILD = ROOT / 'build'
MEDIA = BUILD / 'doc_media'
OUT = ROOT / 'docs' / 'VOID_RUNNER_documentation.docx'

# ---- cover page data (fill in)
COVER = {
    'university': '[University name]',
    'faculty': '[Faculty / Department]',
    'course': 'Digital Systems 2 (Microprocessors)',
    'project': 'Project 1',
    'student': '[Student name]',
    'student_id': '[Student ID]',
    'teacher': '[Teacher name]',
    'semester': '[Semester, e.g. Fall 2026]',
    'date': '[Submission date]',
    'repo': 'https://github.com/Arash1Kazemi/void-runner',
}

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x60, 0x60, 0x60)

# ================================================================= pictures
DUMP_NAMES = [
    'title', 'title_push', 'menu', 'hangar_bolt', 'hangar_dart', 'hangar_twin',
    'opt_lite', 'opt_snd', 'opt_diff', 'opt_test', 'led1', 'led2', 'led3', 'led4',
    'best_empty', 'zone1', 'rockbelt', 'go', 'play1', 'play2', 'pause_score', 'pause_sign',
    'warn1', 'warn2', 'warden1', 'warden2', 'bossdown', 'zonedone', 'pts', 'lifeup',
    'phase2a', 'phase2b', 'hunter1', 'hunter2', 'phase3a', 'phase3b', 'king', 'king2',
    'vic_a', 'vic_b', 'vic_c', 'vic_d', 'over1', 'over2', 'best1', 'best2', 'best3',
]


def screen_png(dump, name, pitch=14):
    path = MEDIA / (name + '.png')
    rd.render([dump], path, pitch=pitch, gap=10)
    return path


def make_media():
    MEDIA.mkdir(parents=True, exist_ok=True)
    sim = BUILD / ('sim.exe' if sys.platform == 'win32' else 'sim')
    log = subprocess.run([str(sim), 'tools/doc_record.script'], cwd=ROOT, capture_output=True,
                         text=True, errors='replace').stdout
    dumps = rd.read_dumps(log)
    assert len(dumps) == len(DUMP_NAMES), 'got %d dumps' % len(dumps)
    for n, d in zip(DUMP_NAMES, dumps):
        if n.startswith('hangar_'):          # the engine flame flickers every 20 ms (screens.h:
            for r in (15, 16):               # frameNo & 1): show the picture with the flame on
                d[3 + r] = d[3 + r][:6] + '###' + d[3 + r][9:]     # (dump rows: 3 bars, then the screen)
    pics = {n: screen_png(d, n) for n, d in zip(DUMP_NAMES, dumps)}
    logo = rd.read_dumps((BUILD / 'doc_logo.txt').read_text(encoding='utf-8'))
    for name, t in [('logo_warp', 600), ('logo_ship', 1600), ('logo_letters', 3100),
                    ('logo_runner', 4200), ('logo_flash', 4900), ('logo_glow', 5400)]:
        pics[name] = screen_png(logo[min(t // 100, len(logo) - 1)], name)
    # busy moments from the clips of docs/media (tools/media_record.script)
    subprocess.run([str(sim), 'tools/media_record.script'], cwd=ROOT, capture_output=True)
    for clip, picks in [('phase1', [44, 120, 200, 300]), ('warden', [66, 108, 148]),
                        ('king', [60, 140, 213])]:
        clipdumps = rd.read_dumps((BUILD / ('rec_%s.txt' % clip)).read_text(encoding='utf-8'))
        for k, i in enumerate(picks):
            pics['%s_%d' % (clip, k)] = screen_png(clipdumps[min(i, len(clipdumps)) - 1],
                                                   '%s_%d' % (clip, k))
    # the whole circuit, lit with a boss-fight picture, and close-ups of its parts
    photo = MEDIA / 'circuit_photo.png'
    subprocess.run([sys.executable, 'tools/render_circuit.py', 'build/rec_warden.txt', '66',
                    str(photo), '2'], cwd=ROOT, check=True)
    pics['circuit'] = photo
    pics.update(circuit_closeups(photo))
    pics['wiring'] = ROOT / 'docs' / 'media' / 'wiring.png'
    pics.update(sprite_sheets())
    return pics


def circuit_closeups(photo):
    """Crops the LED panel, the Mega, the keypad and the buzzers out of the circuit photo."""
    import render_circuit as rc
    gen = rc.gen
    pts = list(gen.PIN_TABLE.values()) + [pt for w in gen.wires for pt in gen.polyline(w)]
    xs, ys = zip(*pts)
    x0, y0, zoom = min(xs) - 80, min(ys) - 80, 2
    im = Image.open(photo)
    boxes = {}
    for kind, key in [('wokwi-led', 'cu_leds'), ('wokwi-pushbutton', 'cu_keypad'),
                      ('wokwi-buzzer', 'cu_buzzers'), ('wokwi-arduino-mega', 'cu_mega')]:
        ps = [p for p in gen.parts if p['type'] == kind]
        size = {'wokwi-led': (40, 50), 'wokwi-pushbutton': (70, 50), 'wokwi-buzzer': (70, 80),
                'wokwi-arduino-mega': (400, 200)}[kind]
        l = min(p['left'] for p in ps) - 30
        t = min(p['top'] for p in ps) - 30
        r = max(p['left'] for p in ps) + size[0] + 30
        b = max(p['top'] for p in ps) + size[1] + 30
        box = tuple(int((v - o) * zoom) for v, o in ((l, x0), (t, y0), (r, x0), (b, y0)))
        path = MEDIA / (key + '.png')
        im.crop(box).save(path)
        boxes[key] = path
    return boxes


def parse_art():
    """The sprites of art.h: {NAME: list of row strings ('#' / '.')}."""
    text = (ROOT / 'void_runner' / 'art.h').read_text(encoding='utf-8')
    art = {}
    for m in re.finditer(r'const uint16_t (\w+)\[\d+\] PROGMEM = \{(.*?)\};', text, re.S):
        rows = [int(b, 2) for b in re.findall(r'0b([01]{16})', m.group(2))]
        if not rows:
            continue
        width = max(16 - (r & -r).bit_length() + 1 for r in rows if r)
        art[m.group(1)] = [''.join('#' if r & (0x8000 >> x) else '.' for x in range(width))
                           for r in rows]
    return art


def draw_sprite(rows, colour, pitch=26):
    w, h = len(rows[0]), len(rows)
    pad = pitch // 2
    im = Image.new('RGB', (w * pitch + 2 * pad, h * pitch + 2 * pad), (40, 40, 40))
    d = ImageDraw.Draw(im)
    for y, row in enumerate(rows):
        for x, c in enumerate(row):
            cx, cy = pad + x * pitch + pitch // 2, pad + y * pitch + pitch // 2
            r = pitch * 5 // 14 if c == '#' else pitch * 4 // 14
            d.ellipse((cx - r, cy - r, cx + r, cy + r), fill=colour if c == '#' else (58, 58, 58))
    return im


def sprite_sheets():
    art = parse_art()
    out = {}
    red, white, green = (255, 70, 70), (240, 240, 240), (70, 255, 90)
    for name, colour in [('SHIP_BOLT', green), ('SHIP_DART', green), ('SHIP_TWIN', green),
                         ('ENEMY_DRONE', red), ('ENEMY_WEAVER', red), ('ENEMY_MINER', red),
                         ('ENEMY_CRUISER', red), ('ROCK_SMALL', white), ('ROCK_BIG', white),
                         ('PICKUP_AMMO', white), ('PICKUP_LIFE', white), ('BOSS_WARDEN', red),
                         ('BOSS_HUNTER', red), ('BOSS_KING', red), ('BOSS_KING_ANGRY', red),
                         ('WARNING_SIGN', white)]:
        path = MEDIA / ('sprite_%s.png' % name.lower())
        draw_sprite(art[name], colour).save(path)
        out['sprite_' + name.lower()] = path
    return out


# ================================================================= document helpers
def shade(paragraph_or_cell, fill):
    el = paragraph_or_cell._p if hasattr(paragraph_or_cell, '_p') else paragraph_or_cell._tc
    pr = el.get_or_add_pPr() if hasattr(el, 'get_or_add_pPr') else el.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pr.append(shd)


class Doc:
    FIG_LABEL, TAB_LABEL = 'Figure', 'Table'          # caption words (the Persian version changes them)

    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)               # A4
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin = sec.bottom_margin = Cm(2)
        normal = self.d.styles['Normal']
        normal.font.name = 'Arial'
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.line_spacing = 1.15
        for lvl, size in [(1, 20), (2, 14), (3, 12)]:
            st = self.d.styles['Heading %d' % lvl]
            st.font.name = 'Arial'
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = ACCENT
            st.paragraph_format.space_before = Pt(14 if lvl > 1 else 0)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
        self.width = Cm(17)
        self.fig = 0
        self.tab = 0

    def h(self, text, level):
        return self.d.add_heading(text, level)

    def p(self, text='', bold_lead=None, size=None, align=None, italic=False, colour=None):
        para = self.d.add_paragraph()
        if bold_lead:
            r = para.add_run(bold_lead)
            r.bold = True
            if size:
                r.font.size = Pt(size)
        r = para.add_run(text)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if colour:
            r.font.color.rgb = colour
        if align:
            para.alignment = align
        return para

    def rich(self, text):
        """A paragraph where **bold** and `code` parts are formatted."""
        para = self.d.add_paragraph()
        for part in re.split(r'(\*\*.+?\*\*|`.+?`)', text):
            if part.startswith('**'):
                para.add_run(part[2:-2]).bold = True
            elif part.startswith('`'):
                r = para.add_run(part[1:-1])
                r.font.name = 'Consolas'
                r.font.size = Pt(9.5)
            elif part:
                para.add_run(part)
        return para

    def bullets(self, items):
        for it in items:
            para = self.rich(it)
            para.style = self.d.styles['List Bullet']
            para.paragraph_format.space_after = Pt(2)

    def caption(self, text):
        self.fig += 1
        para = self.p('%s %d. %s' % (self.FIG_LABEL, self.fig, text), size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
                      italic=True, colour=GREY)
        para.paragraph_format.space_after = Pt(10)

    def picture(self, path, width_cm, caption=None):
        self.d.add_picture(str(path), width=Cm(width_cm))
        para = self.d.paragraphs[-1]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = bool(caption)
        para.paragraph_format.space_after = Pt(2)
        if caption:
            self.caption(caption)

    def gallery(self, items, per_row, width_cm, caption=None):
        """A row of pictures, each with a short label under it (items: (path, label))."""
        rows = (len(items) + per_row - 1) // per_row
        t = self.d.add_table(rows=rows * 2, cols=per_row)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (path, label) in enumerate(items):
            r, c = divmod(i, per_row)
            cell = t.cell(r * 2, c)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.add_run().add_picture(str(path), width=Cm(width_cm))
            lab = t.cell(r * 2 + 1, c).paragraphs[0]
            lab.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = lab.add_run(label)
            run.font.size = Pt(8.5)
            run.font.color.rgb = GREY
        if caption:
            self.caption(caption)
        else:
            self.p('')

    def table(self, header, rows, widths=None, caption=None, size=9.5):
        t = self.d.add_table(rows=1 + len(rows), cols=len(header))
        t.style = self.d.styles['Table Grid']
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, text in enumerate(header):
            cell = t.cell(0, c)
            shade(cell, '1F3A5F')
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for r, row in enumerate(rows, 1):
            for c, text in enumerate(row):
                cell = t.cell(r, c)
                if r % 2 == 0:
                    shade(cell, 'EEF2F7')
                para = cell.paragraphs[0]
                for part in re.split(r'(\*\*.+?\*\*)', str(text)):
                    if not part:
                        continue
                    run = para.add_run(part.strip('*') if part.startswith('**') else part)
                    run.bold = part.startswith('**')
                    run.font.size = Pt(size)
        if widths:
            for row in t.rows:
                for c, w in enumerate(widths):
                    row.cells[c].width = Cm(w)
        if caption:
            self.tab += 1
            para = self.p('%s %d. %s' % (self.TAB_LABEL, self.tab, caption), size=9,
                          align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, colour=GREY)
            para.paragraph_format.space_after = Pt(10)
        else:
            self.p('')
        return t

    def code(self, text, lang):
        """Full source, one paragraph per line: line numbers, comments in green."""
        lines = text.replace('\t', '    ').split('\n')
        if lines and lines[-1] == '':
            lines.pop()
        in_block = False
        digits = len(str(len(lines)))
        for n, line in enumerate(lines, 1):
            para = self.d.add_paragraph()
            pf = para.paragraph_format
            pf.space_after = pf.space_before = Pt(0)
            pf.line_spacing = 1.0
            shade(para, 'F6F8FA')
            num = para.add_run(str(n).rjust(digits) + '  ')
            num.font.name = 'Consolas'
            num.font.size = Pt(7.5)
            num.font.color.rgb = RGBColor(0xA0, 0xA0, 0xA0)
            for chunk, kind in split_code(line, lang, in_block):
                r = para.add_run(chunk)
                r.font.name = 'Consolas'
                r.font.size = Pt(7.5)
                r.font.color.rgb = {'comment': RGBColor(0x2E, 0x7D, 0x32),
                                    'pre': RGBColor(0x8E, 0x24, 0xAA),
                                    'string': RGBColor(0xB2, 0x3A, 0x1D),
                                    'code': RGBColor(0x1A, 0x1A, 0x1A)}[kind]
            in_block = block_state(line, lang, in_block)
        self.p('')

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def block_state(line, lang, in_block):
    if lang != 'c':
        return False
    s = line
    while True:
        if in_block:
            end = s.find('*/')
            if end < 0:
                return True
            s, in_block = s[end + 2:], False
        else:
            s = re.sub(r'"(\\.|[^"\\])*"', '""', s)
            start, lc = s.find('/*'), s.find('//')
            if start < 0 or (0 <= lc < start):
                return False
            s, in_block = s[start + 2:], True


def split_code(line, lang, in_block):
    """Splits one line into (text, kind) parts: comment, pre(processor), string or code."""
    if lang == 'c' and in_block:
        end = line.find('*/')
        if end < 0:
            return [(line, 'comment')]
        return [(line[:end + 2], 'comment')] + split_code(line[end + 2:], lang, False)
    if lang == 'c' and line.lstrip().startswith('#'):
        return [(line, 'pre')]
    comment = '//' if lang == 'c' else '#'
    out, i, start = [], 0, 0
    while i < len(line):
        ch = line[i]
        if ch in '"\'' and not (lang == 'c' and ch == "'" and False):
            j = i + 1
            while j < len(line) and line[j] != ch:
                j += 2 if line[j] == '\\' else 1
            out += [(line[start:i], 'code'), (line[i:j + 1], 'string')]
            i = start = j + 1
            continue
        if line.startswith(comment, i) or (lang == 'c' and line.startswith('/*', i)):
            out.append((line[start:i], 'code'))
            if lang == 'c' and line.startswith('/*', i):
                end = line.find('*/', i + 2)
                if end >= 0:
                    out.append((line[i:end + 2], 'comment'))
                    i = start = end + 2
                    continue
            out.append((line[i:], 'comment'))
            return [p for p in out if p[0]]
        i += 1
    out.append((line[start:], 'code'))
    return [p for p in out if p[0]]


def functions_of(text):
    """(signature, description) of every function: the // comment lines just above it."""
    lines = text.split('\n')
    found = []
    for i, line in enumerate(lines):
        m = re.match(r'^(?!\s)(?:static\s+)?[\w:<>\*&]+(?:\s+[\w:<>\*&]+)*\s+\**(\w+)\(([^)]*)\)\s*\{', line)
        if not m or m.group(1) in ('if', 'for', 'while', 'switch'):
            continue
        desc, k = [], i - 1
        while k >= 0 and lines[k].startswith('//'):
            line_text = lines[k][2:].strip()
            if re.match(r'^[=\-]{3,}', line_text):          # a section divider, not a description
                break
            desc.insert(0, line_text)
            k -= 1
        found.append(('%s(%s)' % (m.group(1), m.group(2).strip()), ' '.join(desc)))
    return found


# ================================================================= the document
def cover(doc):
    d = doc.d
    for _ in range(3):
        doc.p('')
    doc.p(COVER['university'], size=16, align=WD_ALIGN_PARAGRAPH.CENTER).runs[-1].bold = True
    doc.p(COVER['faculty'], size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.p('')
    doc.p(COVER['course'], size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.p(COVER['project'], size=12, align=WD_ALIGN_PARAGRAPH.CENTER, colour=GREY)
    for _ in range(2):
        doc.p('')
    t = doc.p('VOID RUNNER', size=40, align=WD_ALIGN_PARAGRAPH.CENTER)
    t.runs[-1].bold = True
    t.runs[-1].font.color.rgb = ACCENT
    doc.p('A vertical space shooter on an Arduino Mega', size=14,
          align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
    doc.p('432 single LEDs · a 9-button keypad · 4 buzzers · no modules, no libraries', size=10.5,
          align=WD_ALIGN_PARAGRAPH.CENTER, colour=GREY)
    for _ in range(3):
        doc.p('')
    info = [('Student', COVER['student']), ('Student ID', COVER['student_id']),
            ('Teacher', COVER['teacher']), ('Course', COVER['course']),
            ('Semester', COVER['semester']), ('Date', COVER['date']),
            ('Platform', 'Arduino Mega 2560, simulated in Wokwi'), ('Source code', COVER['repo'])]
    t = d.add_table(rows=len(info), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, (k, v) in enumerate(info):
        a, b = t.cell(r, 0), t.cell(r, 1)
        a.width, b.width = Cm(4), Cm(10)
        ra = a.paragraphs[0].add_run(k)
        ra.bold = True
        ra.font.size = Pt(11)
        ra.font.color.rgb = ACCENT
        b.paragraphs[0].add_run(v).font.size = Pt(11)
    doc.page_break()


def hardware_page(doc, pics):
    doc.h('The hardware', 1)
    doc.picture(pics['circuit'], 17, 'The whole circuit in Wokwi: 432 LEDs (left), Arduino Mega 2560 '
                '(top), 3 × 3 push-button keypad and 4 buzzers (right). The LEDs show a real '
                'moment of the game: the WARDEN boss fight.')
    doc.page_break()


def performance(doc, pics):
    P, R = doc.p, doc.rich
    doc.h('3. Project performance', 1)
    R('**VOID RUNNER** is a complete arcade game built only from the parts the course allows: '
      'one Arduino Mega 2560, single LEDs, single push buttons and buzzers. There are no ready-made '
      'modules (no LED matrix module, no keypad module) and no libraries; the screen, the keypad and '
      'the four-voice music are all driven by our own code. The player flies a small ship at the '
      'bottom of a 16 × 24 LED screen, shoots enemies and space rocks, catches ammo crates and '
      'extra lives, and fights a boss at the end of each of the three phases. The project runs in '
      'the Wokwi simulator (Wokwi for VS Code and wokwi.com).')

    doc.h('3.1 Parts and connections', 2)
    doc.table(['Part', 'Count', 'Connection to the Mega'], [
        ['Arduino Mega 2560', '1', '–'],
        ['Single LEDs (wokwi-led), one matrix', '432', '16 × 24 game screen + 3 status bars of 16. '
         '27 row wires (anodes) on pins 22–45 and 47–49, 16 column wires (cathodes) on A0–A15'],
        ['Push buttons (wokwi-pushbutton), 3 × 3 matrix', '9', 'rows on pins 18–20, columns on 15–17'],
        ['Buzzers (wokwi-buzzer)', '4', 'pins 5, 6, 11, 46 (outputs of Timer3, Timer4, Timer1, Timer5)'],
    ], [5, 1.5, 10.5], 'The parts of the device.')
    R('53 of the Mega\'s pins are used, and every header has one job: the double header drives the '
      'LED rows, the analog header the LED columns, the communication header the keypad and the PWM '
      'header the buzzers. No resistors are needed in the simulation. The circuit file '
      '(`diagram.json`) and the pin list of the code (`pins.h`) are both generated by one program '
      '(`tools/gen_circuit.py`), which also checks every wire, so the circuit and the code always '
      'agree.')
    doc.picture(pics['wiring'], 17, 'Simplified wiring: how every LED, key and buzzer is connected.')

    doc.h('3.2 The LED screen', 2)
    R('All 432 LEDs form **one matrix** of 27 rows × 16 columns: the + leg of every LED is on its '
      'row wire, the – leg on its column wire. The Mega switches on **one row at a time** and pulls '
      'the columns of the LEDs that should light LOW. A Timer2 interrupt moves to the next row about '
      '1900 times per second, so the whole screen is redrawn 71 times per second and the eye sees a '
      'steady picture. The columns are written with two port registers (`PORTF`, `PORTK`) at once: '
      'this costs about 3 % of the processor, where `digitalWrite` would cost about 35 %.')
    R('The game draws into a picture in memory (`frame`, one 16-bit word per row) and copies it to '
      'the screen when it is finished, so a half-drawn picture is never shown.')
    doc.table(['Area', 'LED colour', 'Use in the game'], [
        ['Status bar 1', 'red', 'phase progress, then the boss\'s health'],
        ['Status bar 2', 'yellow', 'ammo (blinks when low)'],
        ['Status bar 3', 'green', 'lives (one block of 3 LEDs per life)'],
        ['Game rows 0–5', 'red', 'where enemies and bosses enter, "WARN", "ZONE"'],
        ['Game rows 6–17', 'white', 'open space, message pages'],
        ['Game rows 18–23', 'green', 'the player\'s zone'],
    ], [4, 3, 10], 'LED colours. Each LED has a fixed colour, like the colour strips of the 1978 '
       'Space Invaders cabinet.')
    doc.picture(pics['cu_leds'], 6.5, 'The LED panel in Wokwi during a boss fight.')
    R('**Brightness.** Each LED is lit only 1/27 of the time. Wokwi shows brightness as '
      '(share of time on)^(1/gamma); with the default gamma a lit LED looked only 32 % bright, which '
      'was hard to see. Every LED therefore has gamma = 30 (90 %), which is what a real multiplexed '
      'display does by driving each LED harder in its short time slot. Every LED also has fps = 300, '
      'so moving objects stay sharp, and everything that blinks stays on and off for at least '
      '160 ms so the blink is visible.')
    R('**Text** never scrolls (scrolling letters smeared into a white block in the simulator). '
      'Every message is a still **page** of short words, at most four small letters per line, '
      'shown for 0.9 s, with big digits and signs where a word is not needed.')

    doc.h('3.3 The keypad and the controls', 2)
    R('The 9 push buttons form a 3 × 3 matrix, one button for each key the game uses. The code drives one row LOW at a time and reads the '
      'three columns with the internal pull-ups, so a pressed key reads LOW. The rows that are not '
      'being read are switched to inputs, so two pressed keys can never short a HIGH and a LOW '
      'output together. Every key is debounced. The buttons are labelled in the style of an Atari joystick (UP, DOWN, FIRE ...) and also '
      'react to the PC keyboard in Wokwi.')
    doc.picture(pics['cu_keypad'], 8, 'The keypad in Wokwi.')
    doc.table(['Key', 'Action'], [
        ['W A S D', 'move (tap = one pixel, hold = glide)'],
        ['R', 'fire (tap = one bullet, hold = auto-fire) · OK in the menus'],
        ['Q', 'back · pause in the game'],
        ['P', 'pause / continue'],
        ['M', 'sound on / off (any screen, saved)'],
        ['B', 'brightness 100 / 50 / 20 % (any screen, saved)'],
    ], [3, 14], 'The keys.')
    R('**Firing:** press R for one bullet; keep holding and another follows every 0.2 s (a chirp '
      'when auto-fire starts, and the ship\'s nose blinks). **Moving:** a tap moves one pixel, a '
      'hold makes the ship glide after 0.17 s. The newest key wins: a move key stops the fire, and '
      'while R is held the ship does not glide.')

    doc.h('3.4 The four buzzers', 2)
    R('Each buzzer is connected to the output pin of its own 16-bit hardware timer. In CTC mode with '
      '"toggle on compare match" the timer flips the pin by itself, so a note costs no processor time '
      'and **four notes can sound at once** (Arduino\'s `tone()` can play only one). The timers run '
      'on a slow clock (÷256), which keeps notes up to about C7 within 1 % of their pitch.')
    doc.picture(pics['cu_buzzers'], 8, 'The four buzzers in Wokwi.')
    doc.table(['Buzzer', 'Pin / timer', 'Job'], [
        ['1', '5 / Timer3', 'lead guitar (melody)'],
        ['2', '6 / Timer4', 'bass'],
        ['3', '11 / Timer1', 'every sound effect and menu sound (music only in short jingles)'],
        ['4', '46 / Timer5', 'rhythm guitar: a fifth above the bass (bass + rhythm = a power chord)'],
    ], [2, 3.5, 11.5], 'The voices.')

    doc.h('3.5 Start-up logo and welcome music', 2)
    R('When the device starts, a 5.6-second logo plays in sync with the welcome melody: the picture '
      'runs on the song\'s own clock, so every movement lands on a note. "VOID" is a vertical sign '
      'of bold 10 × 5 letters, and the status bars fill up like a start-up progress bar. Any key '
      'skips the logo.')
    doc.gallery([(pics['logo_warp'], '0.6 s warp'), (pics['logo_ship'], '1.6 s the ship'),
                 (pics['logo_letters'], '3.1 s letters'), (pics['logo_runner'], '4.2 s RUNNER'),
                 (pics['logo_flash'], '4.9 s flash'), (pics['logo_glow'], '5.4 s the sign')],
                6, 2.6, 'The start-up logo.')
    doc.table(['Time', 'Picture', 'Sound'], [
        ['0.0 – 1.2 s', 'warp: star streaks speed up', 'rising whoosh + low A2 drone'],
        ['1.2 – 2.0 s', 'the ship rises and brakes', 'braking notes E4 C4 A3'],
        ['2.0 – 2.4 s', 'the ship shoots up and out', 'quick rising glide'],
        ['2.4 – 3.6 s', 'V, O, I, D light up one by one, flickering on; drones fly past',
         'A4 C5 E5 A5, a crackle on each flicker'],
        ['3.6 – 4.8 s', '"RUNNER" crosses on a dark ribbon, towed by the ship',
         'melody + harmony + bass line'],
        ['4.8 – 5.6 s', 'the whole screen flashes, then the sign glows', 'A-minor chord'],
    ], [2.5, 8, 6.5], 'Timeline of the logo.')

    doc.h('3.6 Title screen, demo and menu', 2)
    R('The title screen keeps the VOID sign lit while ships fly past it, and every 7 seconds a '
      'banner shows the pages "PUSH", a big R, "TO", "PLAY". If nobody touches a key for 20 seconds '
      'a **demo** starts: the autopilot plays the game by itself until any key is pressed. R opens '
      'the menu: PLAY, SHIP, OPTS and BEST (W/S choose, R open, Q back); the chosen item is '
      'underlined.')
    doc.gallery([(pics['title'], 'title'), (pics['title_push'], 'banner'), (pics['menu'], 'menu')],
                3, 3.4, 'Title screen and main menu.')

    doc.h('3.7 Hangar and options', 2)
    R('In the **hangar** the player chooses one of three ships with A/D. The ship is shown three '
      'times bigger, and the three status bars show its speed, ammo and fire power. The **options** '
      'are pages: LITE (brightness), SND (sound), DIFF (EASY / NORM / HARD) and TEST (the LED test). '
      'W/S choose a page, A/D or R change it, page dots at the bottom show where you are. Every '
      'setting is saved in the EEPROM.')
    doc.gallery([(pics['hangar_bolt'], 'BOLT'), (pics['hangar_dart'], 'DART'),
                 (pics['hangar_twin'], 'TWIN')], 3, 3.4, 'The hangar.')
    doc.table(['Ship', 'Speed', 'Fire', 'Ammo', 'Note'], [
        ['BOLT', '70 ms / pixel', 'every 0.20 s', '60', 'balanced'],
        ['DART', '45 ms / pixel', 'every 0.22 s', '40', 'fastest, less ammo'],
        ['TWIN', '95 ms / pixel', 'every 0.26 s', '80', 'two bullets per shot (2 ammo), slowest'],
    ], [2.5, 3, 3, 2, 6.5], 'The three ships.')
    doc.gallery([(pics['opt_lite'], 'LITE'), (pics['opt_snd'], 'SND'), (pics['opt_diff'], 'DIFF'),
                 (pics['opt_test'], 'TEST')], 4, 3, 'The option pages.')

    doc.h('3.8 Playing the game', 2)
    R('A game has three phases. Each phase starts with the pages "ZONE" and a big phase number, the '
      'phase name and "GO!" with a fanfare. Then the enemies and rocks of the phase arrive from its '
      'spawn list. When the list is finished and the screen is empty, "WARN" and a blinking warning '
      'sign appear with a siren while the red bar charges up, and the boss comes with its own music. '
      'When the boss explodes, the pages "ZONE DONE", the bonus points and "LIFE UP" follow, and the '
      'next phase begins. After the last boss the victory screen comes straight away.')
    doc.gallery([(pics['zone1'], 'ZONE 1'), (pics['rockbelt'], 'phase name'), (pics['go'], 'GO!'),
                 (pics['phase1_0'], 'phase 1'), (pics['phase1_1'], 'phase 1'),
                 (pics['phase1_3'], 'phase 1')], 6, 2.6, 'Start of phase 1, ROCK BELT, and play.')
    doc.table(['Topic', 'Rule'], [
        ['Lives', 'EASY 4, NORMAL 3, HARD 3 (at most 4). A hit: explosion, screen shake and 2 s of '
         'blinking shield. +1 life after each phase; extra-life pickups.'],
        ['Ammo', 'Per ship (60 / 40 / 80), full at the start of every phase. Ammo crates give '
         '+20 / 15 / 12 by difficulty. At 5 or less a warning beeps, the yellow bar blinks and a '
         'crate is always sent. Empty: a dry click.'],
        ['Enemies', 'All slower than the player; at most 4 / 5 / 6 on the screen at once (phase '
         '1 / 2 / 3). They shoot only from the upper rows and never right next to the ship, and '
         'they never move onto each other, so the small screen stays readable.'],
        ['Space rocks', 'Hollow, lumpy shapes (small 3 × 3, big 4 × 3) that fall slowly and drift; '
         'a big rock takes 2 hits and breaks into two small ones. They hurt on contact.'],
        ['Explosions', 'Only the bosses and the player\'s ship explode into flying pixels; enemies '
         'just vanish with a boom (flying debris looked like bullets).'],
        ['Collisions', 'Pixel-exact: a hit counts only when lit pixels touch.'],
        ['Score', 'Drone 10, weaver 15, miner 20, cruiser 40, rocks 5 / 10, mine 2, pickup 5, '
         'bosses 200 / 300 / 500, phase bonus 100 × phase; at most 9999.'],
    ], [3, 14], 'The rules of the game.')
    items = [(pics['sprite_enemy_drone'], 'drone'), (pics['sprite_enemy_weaver'], 'weaver'),
             (pics['sprite_enemy_miner'], 'mine-layer'), (pics['sprite_enemy_cruiser'], 'cruiser'),
             (pics['sprite_rock_small'], 'small rock'), (pics['sprite_rock_big'], 'big rock'),
             (pics['sprite_pickup_ammo'], 'ammo crate'), (pics['sprite_pickup_life'], 'extra life')]
    doc.gallery(items, 8, 1.6, 'Enemies, rocks and pickups, pixel by pixel.')
    doc.table(['Enemy', 'Behaviour'], [
        ['Drone', 'sinks slowly, shoots straight down'],
        ['Weaver', 'sinks while zig-zagging, shoots now and then'],
        ['Mine-layer', 'crosses the top of the screen and drops mines (mines can be shot)'],
        ['Cruiser', '5 hits; stops near the top, slides sideways, fires shots aimed at the player, '
         'drops an extra life'],
        ['Rocks', 'fall and drift sideways; a big rock breaks into two'],
    ], [3, 14], 'Enemy types.')

    doc.h('3.9 The status bars', 2)
    R('The three bars above the game screen show the state of the game at a glance. Red: how far '
      'the phase has come; during the warning it charges up, and in the boss fight it shows the '
      'boss\'s health, dropping with every hit. Yellow: the ammo left (it blinks when 5 or fewer are '
      'left). Green: the lives, one block of three LEDs per life.')

    doc.h('3.10 The three phases and their bosses', 2)
    doc.table(['Phase', 'Enemies', 'Boss'], [
        ['1 – ROCK BELT', 'drones and many rocks', '**WARDEN**: wide (10 × 4), slides along the top, '
         '3-way spread, drops rocks; below half health faster with a 5-way spread'],
        ['2 – DEEP MIST', 'weavers, mine-layers, drones', '**HUNTER**: small and fast; dashes to a '
         'new spot, fires an aimed pair, and every 7 s blinks and dives at the player'],
        ['3 – THE VOID', 'cruisers, everything else, rock showers', '**VOID KING**: huge (12 × 5), '
         '5-way spread, calls in drones; at half health its eye opens (roar, shake), then aimed '
         'triple volleys and rock rain'],
    ], [3, 4.5, 9.5], 'Phases and bosses. Boss health: 24 / 22 / 40, scaled by the difficulty.')
    doc.gallery([(pics['warn1'], 'WARN'), (pics['warn2'], 'warning sign'),
                 (pics['warden_0'], 'WARDEN'), (pics['warden_1'], 'WARDEN'),
                 (pics['bossdown'], 'boss down'), (pics['zonedone'], 'ZONE DONE')], 6, 2.6,
                'Phase 1: the warning, the WARDEN fight and the end of the phase.')
    doc.gallery([(pics['pts'], 'bonus'), (pics['lifeup'], 'LIFE UP'), (pics['phase2a'], 'phase 2'),
                 (pics['phase2b'], 'phase 2'), (pics['hunter1'], 'HUNTER'),
                 (pics['hunter2'], 'HUNTER')], 6, 2.6, 'Phase 2, DEEP MIST, and the HUNTER.')
    doc.gallery([(pics['phase3a'], 'phase 3'), (pics['phase3b'], 'phase 3'),
                 (pics['king_0'], 'VOID KING'), (pics['king'], 'VOID KING'),
                 (pics['king_2'], 'eye open'), (pics['king2'], 'form 2')], 6, 2.6,
                'Phase 3, THE VOID, and the VOID KING in both forms.')
    doc.gallery([(pics['sprite_boss_warden'], 'WARDEN'), (pics['sprite_boss_hunter'], 'HUNTER'),
                 (pics['sprite_boss_king'], 'VOID KING'),
                 (pics['sprite_boss_king_angry'], 'VOID KING, form 2')], 4, 3.5,
                'The bosses, pixel by pixel.')
    doc.table(['Difficulty', 'Lives', 'Enemy fire', 'Boss health', 'Ammo crate'], [
        ['EASY', '4', '1.5 × slower', '80 %', '+20'],
        ['NORMAL', '3', 'normal', '100 %', '+15'],
        ['HARD', '3', '1.33 × faster', '130 %', '+12'],
    ], [3, 2, 4, 3.5, 3], 'Difficulty levels (OPTS → DIFF).')

    doc.h('3.11 Pause, game over, victory and best score', 2)
    R('**P** (or Q) pauses the game: the game clock stops, the score and a big pause sign are shown, '
      'and the pages "P:GO" / "Q:EX" explain the keys. When the last life is lost, **GAME OVER** '
      'shows the score (blinking for a new record) and the best score. After the VOID KING the '
      '**victory** screen shows fireworks and the pages "YOU WIN!", "VOID SAFE" and the score. '
      'The **BEST** page shows the high score, how far that game got (ZONE 1–3, or WIN) and its '
      'difficulty; it is kept in the EEPROM, so it survives a power cut on a real Arduino.')
    doc.gallery([(pics['pause_score'], 'pause'), (pics['pause_sign'], 'pause'),
                 (pics['over1'], 'game over'), (pics['over2'], 'game over'),
                 (pics['best_empty'], 'no record yet')], 5, 2.9, 'Pause and game over.')
    doc.gallery([(pics['vic_d'], 'YOU WIN!'), (pics['vic_a'], 'VOID SAFE'), (pics['vic_c'], 'score'),
                 (pics['best1'], 'BEST'), (pics['best2'], 'how far'), (pics['best3'], 'difficulty')],
                6, 2.6, 'Victory and the best score.')

    doc.h('3.12 Sound and music', 2)
    R('All music is written as note lists, one per buzzer, stored in flash. The in-game songs are '
      'original metal tracks written for this project in a short notation '
      '(`tools/compose_music.py`): buzzer 1 plays the lead guitar, buzzer 2 the bass and buzzer 4 a '
      'rhythm guitar a fifth above the bass, which together sound like a power chord. Buzzer 3 is '
      'kept free for sound effects, so an effect never cuts the music.')
    doc.table(['Song', 'Style', 'Character'], [
        ['ROCK BELT (phase 1)', 'fast thrash, E minor, about 190 BPM, loop 10.2 s',
         'the gallop rhythm on low E; a question-and-answer melody'],
        ['DEEP MIST (phase 2)', 'slow and heavy, drop-D, half-time, loop 12.8 s',
         'short muted chugs with gaps; sparse, spooky lead'],
        ['THE VOID (phase 3)', 'syncopated groove, E minor, loop 9.6 s',
         'off-beat stabs; a hook you can hum'],
        ['Boss (all three)', 'tremolo picking, chromatic, about 210 BPM, loop 4.5 s',
         'nonstop 16th notes; a wailing lead figure'],
    ], [4, 6, 7], 'The in-game songs.')
    R('Other music: the start-up melody, the phase fanfare, phase clear, victory and game over. '
      'Sound effects use priorities, so an important sound is never cut by a small one: shot, '
      'auto-fire chirp, empty click, hit, explosion, big explosion, player hit, pickup, extra life, '
      'low ammo, warning siren, boss roar and the LED-test sweep. Explosions are noise (a random '
      'frequency every 6 ms); the other effects slide between two frequencies.')

    doc.h('3.13 LED test', 2)
    R('OPTS → TEST checks the hardware: every LED on, then a row sweep (status bars first), a column '
      'sweep and a checkerboard, with rising beeps. A broken LED or wire is easy to spot.')
    doc.gallery([(pics['led1'], 'all on'), (pics['led2'], 'rows'), (pics['led3'], 'columns'),
                 (pics['led4'], 'checkerboard')], 4, 3, 'The LED test.')

    doc.h('3.14 Reliability and measurements', 2)
    doc.bullets([
        '**Nothing blocks:** there is no `delay()` in the game. A new picture is made every 20 ms, '
        'and each part of the program does one small step per picture. The game has its own clock '
        '(`gameMs`) that stops while paused.',
        '**Screen:** the Timer2 interrupt shows the picture whatever the game code is doing; '
        '3 % of the processor.',
        '**Simulation speed:** the timers run on slow clocks and unused timers are stopped, because '
        'Wokwi spends time on every timer count. Measured: about 93 % of real time during the '
        'logo and 91 % with every LED lit.',
        '**Memory:** only fixed-size arrays (no `malloc`), so nothing can run out during a game. '
        '33 KB of 248 KB flash and 2.0 KB of 8 KB RAM are used; fonts, pixel art, music and spawn '
        'lists are kept in flash.',
        '**Fairness:** at most 4–6 enemies at once, no shots from close range, a crate always '
        'comes when the ammo runs low, and every phase was checked to be winnable.',
    ])

    doc.h('3.15 Testing', 2)
    R('Besides playing it in Wokwi, the real game code is compiled on the PC together with a small '
      'stand-in for the Arduino (registers as variables, a simulated clock, a simulated keypad). '
      'Test scripts press keys and check the Serial output and the screen, so a whole game runs in '
      'under a second. An **autopilot** that plays like a careful player was used to check that '
      'every phase can be won, and to tune the difficulty.')
    doc.table(['Test', 'What it checks', 'Result'], [
        ['controls', 'tap, glide, row limit, single shot, auto-fire rate, move cancels fire, '
         'no glide while firing', 'pass'],
        ['screens', 'logo skip, title, menu, hangar, options, LED test, best, pause, game over and '
         'new record, global keys, power cycle (EEPROM), demo', 'pass'],
        ['playthrough', 'all three phases and bosses, VOID KING form 2, victory', 'pass'],
        ['balance', 'autopilot without shield, reacting only every 250 ms (roughly a human), '
         '48 games per difficulty', 'EASY 47/48, NORMAL 38/48, HARD 27/48 won'],
        ['QA run', 'four test agents played and stress-tested every part (1000+ games and '
         'random key runs)', 'no crash or hang; findings fixed or listed'],
    ], [2.5, 10, 4.5], 'Automatic tests.')
    R('Serial test commands (typed in the Wokwi terminal, starting with `!`) help with testing: '
      '`!d` shows the screen as text, `!s` the game state, `!1`–`!3` start at a phase, `!b` skips '
      'to the boss, `!i` gives a shield and `!x` turns on the autopilot.')


CODE_FILES = [
    ('void_runner/void_runner.ino', 'The main file',
     'The sketch that Wokwi compiles; every other file is a tab of it (one translation unit). '
     '`setup()` stops the unused timers, starts the screen, the keypad and the sound, and loads '
     'the saved settings. `loop()` makes one picture every 20 ms: it reads the keys, handles the '
     'global keys M and B, runs the current screen and plays the sound. The screens form a state '
     'machine (logo → title → menu → game → pause / game over / victory …): each screen function '
     'returns the screen to show next. The Serial test commands (`!d`, `!s`, …) are handled here too.',
     ['menu']),
    ('void_runner/config.h', 'Settings',
     'All settings in one place: the frame time, the names of the keys, the list of screens, the '
     'difficulty tables (lives, enemy fire speed, boss health, ammo per crate) and game constants '
     'such as the number of lives and the phase bonus. Changing the balance of the game only needs '
     'this file.', []),
    ('void_runner/pins.h', 'Pin numbers (generated)',
     'Generated together with the circuit by `tools/gen_circuit.py`, so the pins in the code always '
     'match the wires of the circuit: the 27 LED row pins, the 16 LED column pins (the bits of '
     'PORTF and PORTK), the keypad row and column pins and the four buzzer pins.', ['wiring']),
    ('void_runner/display.h', 'The LED screen',
     'The Timer2 interrupt that lights one LED row at a time (71 full pictures per second), the '
     'picture in memory (`frame`) and every drawing function: pixels, rectangles, bitmaps, letters '
     'of the two fonts, numbers, message pages and the status bars. Brightness (100 / 50 / 20 %) is '
     'made by switching the row off early in its time slot.', ['cu_leds']),
    ('void_runner/input.h', 'The keypad',
     'Scans the 3 × 3 push-button matrix row by row, as taught in the course, and debounces every '
     'key. It remembers which keys are held and which were just pressed or released, so the other '
     'parts can ask "was R just pressed?" or "is D held?".', ['cu_keypad']),
    ('void_runner/sound.h', 'The sound engine',
     'Drives the four buzzers with four 16-bit timers in CTC toggle mode, so each buzzer plays its '
     'note without using processor time. A small music engine plays songs written as note lists '
     '(one per voice, looping or once), and an effect engine on voice 3 plays the sound effects with '
     'priorities. Mute (M) silences everything.', ['cu_buzzers']),
    ('void_runner/tunes.h', 'Songs and sound effects',
     'The note lists of every song and the parts of every sound effect, all stored in flash. The '
     'four in-game songs are generated by `tools/compose_music.py` from a short note notation; the '
     'jingles (fanfare, clear, victory, game over) and the effects are written here directly.', []),
    ('void_runner/storage.h', 'EEPROM: settings and best score',
     'Saves the chosen ship, the difficulty, the brightness, the sound setting and the best score '
     '(with the zone it reached and its difficulty) in the Mega\'s EEPROM, using the EEPROM '
     'registers directly. A check byte tells whether the EEPROM holds valid data.',
     ['best1', 'best3']),
    ('void_runner/font.h', 'Fonts (generated)',
     'Two pixel fonts: small 3 × 5 letters (four fit on a line) and big 5 × 7 letters and digits. '
     'They are drawn as ASCII art in `tools/gen_font.py`, which writes this file.',
     ['zonedone', 'vic_d']),
    ('void_runner/art.h', 'Pixel art',
     'Every picture of the game as binary numbers: one 16-bit word per row, the 1s are the lit '
     'pixels. The three ships, the enemies, the rocks, the pickups, the bosses, the warning sign and '
     'the letters of the logo.',
     ['sprite_ship_bolt', 'sprite_ship_dart', 'sprite_ship_twin', 'sprite_enemy_cruiser',
      'sprite_boss_warden']),
    ('void_runner/logo.h', 'Start-up logo and welcome melody',
     'The 5.6-second logo animation and its melody. The picture is driven by the song\'s own clock, '
     'so every movement lands on a note: warp, the ship, the VOID letters, "RUNNER" on a ribbon, '
     'the flash and the glowing sign.', ['logo_ship', 'logo_letters', 'logo_runner']),
    ('void_runner/title.h', 'Title screen',
     'The VOID sign stays lit, ships keep flying past it, and every 7 seconds a banner shows the '
     'pages "PUSH", a big R, "TO", "PLAY". After 20 s without a key the demo starts.',
     ['title', 'title_push']),
    ('void_runner/menu.h', 'Main menu',
     'The menu PLAY, SHIP, OPTS, BEST: W/S move the underline, R opens the item, Q goes back to '
     'the title.', ['menu']),
    ('void_runner/world.h', 'Everything on the game screen',
     'The data of the game world in fixed-size pools: the player, bullets and mines, enemies and '
     'rocks, pickups, explosion particles and the boss, with the tables of the three ships and the '
     'enemy types and small helpers to create them. There is no `malloc`, so nothing can run out '
     'during a game.', ['phase1_1']),
    ('void_runner/player.h', 'The player\'s ship',
     'The key rules for moving and firing (tap, hold, auto-fire, newest key wins), ammo and the '
     'low-ammo warning, getting hit (a life, explosion, shake, shield) and drawing the ship.',
     ['hangar_bolt', 'hangar_dart', 'hangar_twin']),
    ('void_runner/enemies.h', 'Enemies and rocks',
     'How each enemy type moves and shoots, and the rule that keeps them apart: a move may not bring '
     'an enemy closer than one empty pixel to another one, so it steps aside, turns or waits. New '
     'enemies take the nearest free column.',
     ['sprite_enemy_drone', 'sprite_enemy_weaver', 'sprite_enemy_miner', 'sprite_rock_small',
      'sprite_rock_big']),
    ('void_runner/levels.h', 'The three phases',
     'The spawn list of each phase (wait, what, where), its enemy limit, its name, its song and its '
     'boss. The lists are stored in flash.', ['zone1', 'rockbelt', 'go']),
    ('void_runner/boss.h', 'The bosses',
     'The behaviour of the WARDEN, the HUNTER (a small state machine: dash, aim, charge) and the '
     'VOID KING (two forms), their health and their drawing.',
     ['warden_1', 'hunter1', 'king2']),
    ('void_runner/game.h', 'One game',
     'Runs one game: the stages of a phase (intro, waves, warning, boss, boss down, clear), the '
     'spawner, all collisions (bullets, enemies, rocks, mines, pickups, boss), the score, the '
     'messages on top of the game and the three status bars.', ['warn2', 'bossdown', 'pts']),
    ('void_runner/autopilot.h', 'Autopilot',
     'A bot that plays like a careful player. It is used by the title screen\'s demo and by the '
     'automatic tests: for every column it works out how soon something would hit the ship there, '
     'flies to the safest column near its target and fires when lined up.', ['phase1_3']),
    ('void_runner/screens.h', 'The other screens',
     'The hangar, the options (with the LED test), the best score, pause, game over and victory. '
     'Each screen function handles its keys, draws one picture and returns the next screen.',
     ['opt_diff', 'pause_sign', 'over1', 'led4']),
]

TOOL_FILES = [
    ('tools/gen_circuit.py', 'Circuit generator',
     'Generates `diagram.json` (432 LEDs, 9 buttons, 4 buzzers and every wire) and `pins.h`, '
     'routes the wires so they never cross, and checks every connection.'),
    ('tools/gen_font.py', 'Font generator', 'Turns the ASCII-art letters into `font.h`.'),
    ('tools/compose_music.py', 'Music composer',
     'The in-game songs in a short note notation; writes their note lists into `tunes.h` and '
     'can render them as WAV files to listen to.'),
    ('tools/render_dumps.py', 'Screen renderer',
     'Draws screen dumps (the `!d` command) as PNG pictures or GIF animations in the real LED colours.'),
    ('tools/draw_wiring.py', 'Wiring diagram', 'Draws the simplified wiring diagram.'),
    ('tools/render_circuit.py', 'Circuit photo',
     'Draws the whole circuit with Wokwi\'s own part pictures, lit with a game picture.'),
    ('tools/make_doc.py', 'This document', 'Builds this document from the code and the game itself.'),
    ('tests/host/Arduino.h', 'PC stand-in for Arduino',
     'Registers as variables, a simulated clock and keypad, so the real game code runs on a PC.'),
    ('tests/host/sim.cpp', 'PC test runner', 'Runs the game code from test scripts (keys, Serial, dumps).'),
    ('tests/host/run_tests.py', 'Test runner', 'Builds the PC version and runs all tests and the balance check.'),
    ('tests/host/controls.txt', 'Test: controls', 'The key rules.'),
    ('tests/host/screens.txt', 'Test: screens', 'Every screen and the EEPROM.'),
    ('tests/host/playthrough.txt', 'Test: whole game', 'All phases to the victory.'),
]


def code_part(doc, pics):
    doc.page_break()
    doc.h('4. Project code', 1)
    R = doc.rich
    R('The program is written in C++ for the Arduino Mega (Arduino core functions plus direct '
      'register access where speed or timing needs it; no libraries). It is one sketch, '
      '`void_runner.ino`, with one header file per module; in Wokwi every file is a tab. All files '
      'together are compiled as one unit. The game logic advances in steps of 20 ms; the screen and '
      'the sound run on hardware timers.')
    rows = []
    for path, title, _, _ in CODE_FILES:
        n = len((ROOT / path).read_text(encoding='utf-8').splitlines())
        rows.append([Path(path).name, title, str(n)])
    total = sum(int(r[2]) for r in rows)
    rows.append(['**total**', '', '**%d**' % total])
    doc.table(['File', 'Contents', 'Lines'], rows, [4.5, 10.5, 2], 'The files of the program.')
    R('Each file below starts with a short explanation and the list of its functions (taken from the '
      'description written above every function), followed by the full source code.')
    for i, (path, title, text, pic_names) in enumerate(CODE_FILES, 1):
        src = (ROOT / path).read_text(encoding='utf-8')
        doc.page_break()
        doc.h('4.%d %s – %s' % (i, Path(path).name, title), 2)
        R(text)
        if pic_names:
            per = len(pic_names)
            width = {1: 3.6, 2: 3.4, 3: 3.2, 4: 3, 5: 2.8}[per]
            if pic_names[0] in ('wiring', 'cu_leds', 'cu_keypad', 'cu_buzzers'):
                doc.picture(pics[pic_names[0]], {'wiring': 15, 'cu_leds': 5}.get(pic_names[0], 7),
                            'Belongs to %s.' % Path(path).name)
            else:
                doc.gallery([(pics[n], '') for n in pic_names], per, width,
                            'What %s draws or uses.' % Path(path).name)
        funcs = functions_of(src)
        if funcs:
            doc.table(['Function', 'What it does'], [[f, d or '–'] for f, d in funcs], [6, 11],
                      size=8.5)
        doc.h('Source code', 3)
        doc.code(src, 'c')

    doc.page_break()
    doc.h('4.%d Development tools and tests' % (len(CODE_FILES) + 1), 2)
    R('These programs are not part of the Arduino sketch. They generate parts of it (the circuit, '
      'the pin list, the font, the music), make the pictures, and test the game on the PC.')
    doc.table(['File', 'Job'], [[p, t + ': ' + e] for p, t, e in TOOL_FILES], [5, 12], size=9)
    for path, title, text in TOOL_FILES:
        src = (ROOT / path).read_text(encoding='utf-8')
        doc.h('%s – %s' % (path, title), 3)
        R(text)
        doc.code(src, 'c' if path.endswith(('.h', '.cpp')) else 'py')


def main():
    pics = make_media()
    doc = Doc()
    cover(doc)
    hardware_page(doc, pics)
    performance(doc, pics)
    code_part(doc, pics)
    OUT.parent.mkdir(exist_ok=True)
    doc.d.save(OUT)
    print('wrote', OUT)


if __name__ == '__main__':
    main()
